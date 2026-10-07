"""Build a concise PDF and PowerPoint summary of the calibrated optical model."""

from __future__ import annotations

import tempfile
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyArrowPatch, Rectangle
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from forward_model import ForwardModel, render_display_canvas


OUT = Path(__file__).resolve().parent
W, H = 13.333, 7.5
INK = "#111111"
GRAY = "#555555"
LIGHT = "#dddddd"
SCENE_MAP = {
    "dusk": "scene_08_lake_dusk",
    "citystreet": "scene_03_citystreet",
    "falls": "scene_11_multnomah_falls",
    "forest": "scene_19_snow_forest",
    "harbor": "scene_14_city_harbor",
    "london": "scene_04_busy_london",
    "mountain": "scene_02_mountain",
    "produce": "scene_15_market_produce",
    "rocks": "scene_10_waterfall_rocks",
    "skyline": "scene_13_london_skyline",
    "street": "scene_05_prague_street",
    "vineyard": "scene_18_sonoma_vineyard",
}
SCENE_CREDIT = {
    "dusk": "C. Stone",
    "citystreet": "Artem Mihailov",
    "falls": "Chris Briggs",
    "forest": "Fabian Kleiser",
    "harbor": "Lukas Rychvalsky",
    "london": "Dust Studio",
    "mountain": "Aleks Dahlberg",
    "produce": "Maria Wang",
    "rocks": "Gabe Hobbs",
    "skyline": "Rowan Freeman",
    "street": "Alice",
    "vineyard": "Trent Erwin",
}


def header(fig, title: str, subtitle: str, page: int) -> None:
    fig.text(0.045, 0.955, title, ha="left", va="top", fontsize=23,
             fontweight="bold", color=INK)
    fig.text(0.047, 0.908, subtitle, ha="left", va="top", fontsize=10.5,
             color=GRAY)
    fig.add_artist(plt.Line2D([0.045, 0.955], [0.885, 0.885],
                              transform=fig.transFigure, color=INK, lw=0.8))
    fig.text(0.955, 0.022, f"{page} / 6", ha="right", va="bottom",
             fontsize=9, color=GRAY)


def new_slide(title: str, subtitle: str, page: int):
    fig = plt.figure(figsize=(W, H), facecolor="white")
    header(fig, title, subtitle, page)
    return fig


def image_panel(ax, image: np.ndarray, title: str | None = None,
                vmin=None, vmax=None, cmap="gray", aspect="equal") -> None:
    ax.imshow(image, cmap=cmap, vmin=vmin, vmax=vmax,
              interpolation="nearest", aspect=aspect)
    ax.set_axis_off()
    if title:
        ax.set_title(title, fontsize=11, pad=5, color=INK, fontweight="bold")


def add_axes(fig, x, y, w, h):
    return fig.add_axes([x, y, w, h])


def green_canvas(path: Path, model: ForwardModel) -> np.ndarray:
    with Image.open(path) as image:
        rgb = np.asarray(image.convert("RGB"), dtype=np.uint8)
    return render_display_canvas(rgb, model.display_shape, scale=0.9, channel="green")


def center_square(image: np.ndarray, side: int = 972) -> np.ndarray:
    y, x = image.shape[:2]
    y0, x0 = (y - side) // 2, (x - side) // 2
    return image[y0:y0 + side, x0:x0 + side]


def scene_reference(label: str, model: ForwardModel) -> tuple[np.ndarray, Path]:
    source = ROOT / "assets_scenes" / "square_1600" / (SCENE_MAP[label] + ".png")
    return green_canvas(source, model), source


def metrics(actual: np.ndarray, predicted: np.ndarray, dark: np.ndarray) -> tuple[float, float, float]:
    error = predicted - actual.astype(np.float32)
    rmse = float(np.sqrt(np.mean(error ** 2)))
    correlation = float(np.corrcoef((actual.astype(np.float32) - dark).ravel(),
                                    (predicted - dark).ravel())[0, 1])
    relative = rmse / max(float(np.mean(actual.astype(np.float32) - dark)), 1.0)
    return rmse, correlation, relative


def slide1(model: ForwardModel):
    fig = new_slide(
        "A compact model from display pixels to RAW16",
        "A linear-intensity approximation for the green display channel and a fixed lenslet camera setup",
        1,
    )
    nodes = [
        ("1600 × 1600\nscene / letter", 0.145),
        ("Green only\ncentered at 90%", 0.145),
        ("Measured display\nresponse R[u]", 0.155),
        ("PSF convolution\n+ 10% spot copies", 0.175),
        ("Flat gain × scale\n+ black background", 0.17),
        ("Predicted\nRAW16 frame", 0.13),
    ]
    gap = 0.014
    total = sum(width for _, width in nodes) + gap * (len(nodes) - 1)
    x = (1 - total) / 2
    y, height = 0.64, 0.14
    for index, (label, width) in enumerate(nodes):
        fig.add_artist(Rectangle((x, y), width, height, transform=fig.transFigure,
                                 facecolor="white", edgecolor=INK, linewidth=1.0))
        fig.text(x + width / 2, y + height / 2, label, ha="center", va="center",
                 fontsize=9, color=INK)
        if index < len(nodes) - 1:
            arrow = FancyArrowPatch((x + width + 0.002, y + height / 2),
                                    (x + width + gap - 0.002, y + height / 2),
                                    transform=fig.transFigure, arrowstyle="->",
                                    mutation_scale=10, linewidth=0.9, color=INK)
            fig.add_artist(arrow)
        x += width + gap

    fig.text(0.5, 0.53,
             r"$\hat y = d + s\,g\odot\{0.90\,C_h(R[u]) + 0.10\,L(R[u])\}$",
             ha="center", va="center", fontsize=21, color=INK)
    fig.text(0.5, 0.46,
             r"$d$: monitor-black map   ·   $g$: white-field pixel gain   ·   $h$: measured 2-D PSF   ·   $L$: weighted lenslet copies",
             ha="center", va="center", fontsize=10.5, color=GRAY)

    # Concise acquisition facts and model scope.
    cards = [
        (0.06, "DISPLAY", "1920 × 1080\nGreen channel; 90% fullscreen"),
        (0.29, "CAMERA", "ASI174MM · RAW16\n80 ms · gain 0"),
        (0.52, "CALIBRATION", "120 spot centers\nMeasured response, background, flat field"),
        (0.75, "ASSUMPTION", "Incoherent light\nIntensity contributions add linearly"),
    ]
    for x, label, detail in cards:
        fig.text(x, 0.34, label, fontsize=9, fontweight="bold", color=INK)
        fig.text(x, 0.30, detail, fontsize=10, color=INK, va="top", linespacing=1.35)
    fig.text(0.06, 0.14,
             "Model terms are shared across the alphabet and 20 square scenes; scene references are used only in the evaluation figures.",
             fontsize=10, color=GRAY)
    fig.text(0.06, 0.095,
             f"Fitted constants: signal scale {model.signal_scale:.3f} · PSF sampling {model.psf_sampling_xy[0]:.2f}× / {model.psf_sampling_xy[1]:.2f}× · spot mix {model.spot_mix:.2f}",
             fontsize=10, color=INK)
    return fig


def slide2(model: ForwardModel):
    fig = new_slide(
        "What is fitted in the forward model?",
        "Each map/curve comes from calibration captures; detector maps are shown with robust contrast limits.",
        2,
    )
    dark = model.dark
    gain = model.flat_gain
    psf_view = np.sqrt(np.maximum(model.psf, 0) / max(float(model.psf.max()), 1e-12))
    panels = [
        (0.055, 0.54, 0.26, 0.27, "1  Monitor-black background", dark,
         np.percentile(dark, 1), np.percentile(dark, 99), "RAW16 DN; includes ambient light"),
        (0.37, 0.54, 0.26, 0.27, "2  Flat-field pixel gain", gain,
         np.percentile(gain, 1), np.percentile(gain, 99), "White field − background"),
        (0.685, 0.54, 0.26, 0.27, "3  Measured 2-D PSF", psf_view,
         0, 1, "Square-root display reveals weak spots"),
    ]
    for x, y, w, h, title, image, vmin, vmax, caption in panels:
        ax = add_axes(fig, x, y, w, h)
        image_panel(ax, image, title, vmin=vmin, vmax=vmax)
        fig.text(x + w / 2, y - 0.035, caption, ha="center", va="top",
                 fontsize=8.5, color=GRAY)

    ax = add_axes(fig, 0.075, 0.17, 0.36, 0.25)
    codes = np.arange(256)
    ax.plot(codes, model.response, color=INK, lw=2)
    ax.scatter(codes[::32], model.response[::32], color=INK, s=12, zorder=3)
    ax.set_xlim(0, 255)
    ax.set_ylim(0, 1.03)
    ax.set_xlabel("Green display code", fontsize=9)
    ax.set_ylabel("Normalized linear intensity", fontsize=9)
    ax.set_title("4  Display response lookup table", fontsize=11, pad=5,
                 color=INK, fontweight="bold")
    ax.grid(color="#dddddd", linewidth=0.5)
    ax.tick_params(labelsize=8, colors=INK)

    ax = add_axes(fig, 0.56, 0.17, 0.36, 0.25)
    weights = np.asarray(model.spot_weights, dtype=np.float32)
    sizes = 8 + 45 * weights / max(float(weights.max()), 1e-6)
    ax.scatter(model.spots_yx[:, 1], model.spots_yx[:, 0], s=sizes,
               c=weights, cmap="Greys", edgecolors=INK, linewidths=0.3)
    ax.set_xlim(0, model.dark.shape[1])
    ax.set_ylim(model.dark.shape[0], 0)
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel("Detector x (pixels)", fontsize=9)
    ax.set_ylabel("Detector y (pixels)", fontsize=9)
    ax.set_title(f"5  Weighted lenslet centers ({len(model.spots_yx)})",
                 fontsize=11, pad=5, color=INK, fontweight="bold")
    ax.tick_params(labelsize=8, colors=INK)
    fig.text(0.5, 0.085,
             f"Spot scale |mₓ|={abs(model.magnification_xy[0]):.3f}, |mᵧ|={abs(model.magnification_xy[1]):.3f}  ·  PSF sampling x={model.psf_sampling_xy[0]:.2f}, y={model.psf_sampling_xy[1]:.2f}  ·  scalar signal scale={model.signal_scale:.3f}",
             ha="center", fontsize=9.3, color=INK)
    return fig


def slide3(model: ForwardModel):
    fig = new_slide(
        "Forward prediction against measured RAW16",
        "The two rows show an asymmetric letter and a textured scene; residual = prediction − measurement.",
        3,
    )
    letter = ROOT / "assets_letters" / "02_B.png"
    scene_ref, scene_path = scene_reference("dusk", model)
    examples = [
        ("Letter B", green_canvas(letter, model),
         np.load(ROOT / "forward_model_validation" / "letters" / "measurements" / "B_actual_raw16.npy")),
        ("Scene: dusk", scene_ref,
         np.load(ROOT / "forward_model_validation" / "scenes" / "square_letter_style" / "dusk_actual_raw16.npy")),
    ]
    col_x = [0.10, 0.36, 0.62]
    col_titles = ["Measured signal", "Model prediction", "Residual"]
    for x, title in zip(col_x, col_titles):
        fig.text(x + 0.115, 0.855, title, ha="center", fontsize=11,
                 fontweight="bold", color=INK)
    for row, (name, canvas, actual) in enumerate(examples):
        pred = model.predict(canvas)
        rmse, corr, rel = metrics(actual, pred, model.dark)
        actual_signal = np.maximum(actual.astype(np.float32) - model.dark, 0)
        pred_signal = np.maximum(pred - model.dark, 0)
        vmax = float(np.percentile(np.concatenate((actual_signal.ravel(),
                                                   pred_signal.ravel())), 99.7))
        residual = pred - actual.astype(np.float32)
        lim = max(float(np.percentile(np.abs(residual), 99.5)), 1.0)
        y = 0.535 if row == 0 else 0.185
        h = 0.265
        fig.text(0.047, y + h / 2, name, rotation=90, ha="center", va="center",
                 fontsize=11, fontweight="bold", color=INK)
        ax = add_axes(fig, col_x[0], y, 0.23, h)
        image_panel(ax, actual_signal, vmin=0, vmax=vmax)
        ax = add_axes(fig, col_x[1], y, 0.23, h)
        image_panel(ax, pred_signal, vmin=0, vmax=vmax)
        ax = add_axes(fig, col_x[2], y, 0.23, h)
        image_panel(ax, residual, vmin=-lim, vmax=lim, cmap="coolwarm")
        fig.text(0.87, y + h / 2,
                 f"RMSE\n{rmse:.0f} DN\nr {corr:.3f}\nrel. {rel:.2f}",
                 ha="left", va="center", fontsize=8.2, color=INK, linespacing=1.35)
    fig.text(0.1, 0.105,
             "Signal panels subtract the fitted monitor-black map; both use a shared scale within each row.",
             fontsize=9, color=GRAY)
    fig.text(0.1, 0.073,
             "Residual panels use a symmetric scale; RAW16 measurements are at 80 ms exposure and gain 0.",
             fontsize=9, color=GRAY)
    return fig


def slide4(model: ForwardModel):
    fig = new_slide(
        "Measurement-only reconstruction: letters",
        "A, B, F, and O reconstructed by flat-field correction, spot-cell folding, and median averaging.",
        4,
    )
    letters = ["A", "B", "F", "O"]
    side = 972
    for i, letter in enumerate(letters):
        col, row = i % 2, i // 2
        x = 0.055 + col * 0.48
        y = 0.505 if row == 0 else 0.135
        fig.text(x + 0.005, y + 0.25, f"Letter {letter}", fontsize=13,
                 fontweight="bold", color=INK)
        asset = sorted((ROOT / "assets_letters").glob(f"*_{letter}.png"))[0]
        reference = center_square(green_canvas(asset, model), side)
        recon_path = ROOT / "forward_model_validation" / "letters" / "reconstructions" / f"{letter}_spot_reconstruction.png"
        with Image.open(recon_path) as im:
            recon = center_square(np.asarray(im.convert("L")), side)
        for px, image, label in [(x, reference, "Displayed"),
                                 (x + 0.22, recon, "Reconstructed")]:
            ax = add_axes(fig, px, y + 0.035, 0.205, 0.205)
            image_panel(ax, image, aspect="equal")
            fig.text(px + 0.1025, y + 0.012, label, ha="center", fontsize=8.5,
                     color=GRAY)
    fig.text(0.055, 0.075,
             "Each reconstruction uses only its RAW16 measurement (120 copies); the displayed image is included only for visual comparison.",
             fontsize=9, color=GRAY)
    return fig


def scene_slide(model: ForwardModel, labels: list[str], page: int, title: str):
    fig = new_slide(
        title,
        "Green-only display at 90%; reference and reconstruction are shown as centered square crops.",
        page,
    )
    side = 972
    for i, label in enumerate(labels):
        col, row = i % 3, i // 3
        x = 0.045 + col * 0.308
        y = 0.515 if row == 0 else 0.16
        ref, _ = scene_reference(label, model)
        recon_path = ROOT / "forward_model_validation" / "scenes" / "square_letter_style" / "reconstructions" / f"{label}_spot_reconstruction.png"
        with Image.open(recon_path) as im:
            recon = center_square(np.asarray(im.convert("L")), side)
        fig.text(x + 0.002, y + 0.205, label.title(), fontsize=11,
                 fontweight="bold", color=INK)
        for dx, image, tag in [(0, center_square(ref, side), "Displayed"),
                               (0.151, recon, "Reconstructed")]:
            ax = add_axes(fig, x + dx, y + 0.025, 0.144, 0.144)
            image_panel(ax, image)
            fig.text(x + dx + 0.072, y + 0.006, tag, ha="center",
                     fontsize=7.5, color=GRAY)
    credits = " · ".join(f"{label.title()}: {SCENE_CREDIT[label]}" for label in labels)
    fig.text(0.047, 0.078,
             "Spot-fold reconstructions use RAW16 measurements only; references are for comparison. " + credits,
             fontsize=7.3, color=GRAY)
    return fig


def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    model = ForwardModel.load(ROOT / "forward_model.npz")
    slides = [
        slide1(model),
        slide2(model),
        slide3(model),
        slide4(model),
        scene_slide(model, ["dusk", "citystreet", "falls", "forest", "harbor", "london"],
                    5, "Measurement-only reconstruction: scenes (1/2)"),
        scene_slide(model, ["mountain", "produce", "rocks", "skyline", "street", "vineyard"],
                    6, "Measurement-only reconstruction: scenes (2/2)"),
    ]
    pdf_path = OUT / "forward_model_overview.pdf"
    pptx_path = OUT / "forward_model_overview.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(W)
    presentation.slide_height = Inches(H)
    blank = presentation.slide_layouts[6]
    with tempfile.TemporaryDirectory(prefix="forward_model_slides_") as tmp, PdfPages(pdf_path) as pdf:
        for index, fig in enumerate(slides, start=1):
            png_path = Path(tmp) / f"slide_{index:02d}.png"
            fig.savefig(png_path, dpi=150, facecolor="white")
            pdf.savefig(fig, facecolor="white")
            slide = presentation.slides.add_slide(blank)
            slide.shapes.add_picture(str(png_path), 0, 0,
                                     width=presentation.slide_width,
                                     height=presentation.slide_height)
            plt.close(fig)
    presentation.core_properties.title = "Compact optical display-to-camera forward model"
    presentation.core_properties.subject = "Calibrated RAW16 forward model and measurement-only reconstructions"
    presentation.save(pptx_path)
    print(f"Saved {pptx_path}")
    print(f"Saved {pdf_path}")
    print(f"Slides: {len(presentation.slides)}")


if __name__ == "__main__":
    build()
